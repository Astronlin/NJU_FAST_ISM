"""Bounded reference reads and streaming telemetry; retained reference smoothing."""
from pathlib import Path
import hashlib
import numpy as np
from scipy.ndimage import uniform_filter1d
from scipy.interpolate import interp1d
from astropy.time import Time
import astropy.units as u
from fastfish.io.fits_stream import fixed_table
def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()

def bracket(x, value, right=False):
    """Binary search without copying a huge strided FITS time column."""
    lo,hi=0,len(x)
    while lo<hi:
        mid=(lo+hi)//2
        if float(x[mid])<value or (right and float(x[mid])==value):lo=mid+1
        else:hi=mid
    return lo


def smoothing_rows(n_t=10):
    value=int(n_t)
    if value!=n_t or not 1<=value<=3000:
        raise ValueError('n_t must be an integer from 1 to 3000; 1 disables time averaging')
    return value


def reference_data(path, tmin, tmax, science_freq, verify=False, n_t=10):
    n_t=smoothing_rows(n_t)
    with fixed_table(path,extension=2,random_access=True) as (tab,h):
        times=tab['MJD'];n=len(times)
        if tmin<float(times[0]) or tmax>float(times[-1]):
            raise ValueError('Selected science spectra extend beyond temperature reference')
        # Enough context for the chosen filter; retain a useful diagnostic time span.
        pad=max(160,(n_t+1)//2+2)
        i0=max(0,bracket(times,tmin)-pad);i1=min(n,bracket(times,tmax,True)+pad)
        nf=int(h[0].header['OBSNCHAN']);df=float(h[1].header['CHAN_BW'])/1e3
        cf=(float(h[0].header['OBSFREQ'])-float(h[0].header['OBSBW'])/2)/1e3+np.arange(nf)*df
        j0=max(0,np.searchsorted(cf,science_freq[0])-2)
        j1=min(nf,np.searchsorted(cf,science_freq[-1])+3)
        temp=tab['TSYS_PSR'][i0:i1,:2,j0:j1].astype(float)
        tt=times[i0:i1].copy();cf=cf[j0:j1]
        if not np.isfinite(tt).all() or np.any(np.diff(tt)<=0):
            raise ValueError('Temperature-reference times must increase')
        finite=np.isfinite(temp)
        # The finite case is exactly the historical filter. Normalization keeps
        # isolated NaNs from poisoning a complete time suffix of a channel.
        if n_t==1:smoothed=np.where(finite,temp,np.nan)
        elif finite.all():smoothed=uniform_filter1d(temp,n_t,axis=0,mode='nearest')
        else:
            support=uniform_filter1d(finite.astype(float),n_t,axis=0,mode='nearest')
            total=uniform_filter1d(np.where(finite,temp,0.),n_t,axis=0,mode='nearest')
            smoothed=np.divide(total,support,out=np.full_like(total,np.nan),where=support>=.9)
        # Three finite channels independently verify crop+padding against the
        # corresponding full-time historical smoothing on the same channels.
        errors=[]
        for local_j in (sorted(set([0,len(cf)//2,len(cf)-1])) if verify else []):
            full=tab['TSYS_PSR'][:,0,j0+local_j].astype(float)
            if np.isfinite(full).all():
                historical=uniform_filter1d(full,n_t,mode='nearest')
                select=(tt>=tmin)&(tt<=tmax)
                error=float(np.max(np.abs(historical[i0:i1][select]-smoothed[:,0,local_j][select])))
                errors.append(error)
        info={'temperature_smoothing_rows':n_t,'rows_read':len(tt),'full_rows':n,'frequency_channels_read':len(cf),'full_frequency_channels':nf,
              'time_slice':[i0,i1],'frequency_slice':[int(j0),int(j1)],
              'nonfinite_input_fraction':float((~finite).mean()),'minimum_smoothing_support':.9,
              'finite_channel_full_filter_comparison_max_K':max(errors) if errors else None}
    return dict(mjd=tt, frequency_ghz=cf, raw=temp, smoothed=smoothed, info=info)


def reference_slice(path, tmin, tmax, science_freq, verify=False, n_t=10):
    pack=reference_data(path,tmin,tmax,science_freq,verify=verify,n_t=n_t)
    return interp1d(pack["mjd"],pack["smoothed"],axis=0,bounds_error=True),pack["frequency_ghz"],pack["info"]


def telemetry_cache(scan, sheet, tmin, tmax, cache_dir):
    from openpyxl import load_workbook
    out=Path(cache_dir)/(scan+'.npz');out.parent.mkdir(parents=True,exist_ok=True)
    lower=(Time(tmin,format='mjd')+8*u.hr-5*u.s).iso
    upper=(Time(tmax,format='mjd')+8*u.hr+5*u.s).iso
    keys=['SysTime','SDP_PhaPos_X','SDP_PhaPos_Y','SDP_PhaPos_Z','SDP_AngleM']
    book=load_workbook(sheet,read_only=True,data_only=True)
    try:
        rows=book.active.iter_rows(values_only=True);head=next(rows)
        columns=[head.index(k) for k in keys];data={k:[] for k in keys}
        for row in rows:
            if row[columns[0]] is None:continue
            value=str(row[columns[0]]).replace('T',' ')
            if value<lower:continue
            if value>upper:break
            for k,i in zip(keys,columns):data[k].append(value if k=='SysTime' else row[i])
    finally:book.close()
    if len(data['SysTime'])<2:raise ValueError('No matching telemetry: '+scan)
    arrays={k:np.asarray(v,dtype=str if k=='SysTime' else float) for k,v in data.items()}
    np.savez_compressed(out,**arrays)
    print('TELEMETRY',scan,len(data['SysTime']),flush=True)
    return {'path':str(out),'rows':len(data['SysTime']),'sha256':sha(out),'source':str(sheet),
            'source_bytes':Path(sheet).stat().st_size,'system_time_first_last':[data['SysTime'][0],data['SysTime'][-1]]}


